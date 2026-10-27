import os
import docx

doc_path = r"C:\Users\user\PycharmProjects\the_haldankar_method_research_laboratory\BusinessQuartelyFirst\Transformed_HDFC_Bank_Q1_2026_27_Financials.docx"

if not os.path.exists(doc_path):
    print("Error: Document file does not exist!")
    exit(1)

doc = docx.Document(doc_path)

print(f"Document Size: {os.path.getsize(doc_path)} bytes")
print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
print(f"Total Headings: {len(headings)}")

# Check key data points
full_text = "\n".join([p.text for p in doc.paragraphs])
for t in doc.tables:
    for row in t.rows:
        full_text += "\n" + " | ".join([cell.text for cell in row.cells])

check_terms = [
    "79,362.78", "19,059.72", "35,846.35", "1.17%", "0.41%", "19.57%",
    "7,67,70,39,761", "13,51,35,135", "740", "1,144.46", "Atanu Chakraborty",
    "Dubai Financial Services Authority", "440", "20,785.14", "90,575.33",
    "19,244.71", "HDFC Life Insurance Company Limited", "B S R & Co. LLP",
    "Batliboi & Purohit", "26041870OATJRT2952", "26116976GFKHWQ9629"
]

missing = []
for term in check_terms:
    if term not in full_text:
        missing.append(term)

if missing:
    print(f"Warning: The following expected terms were NOT found: {missing}")
else:
    print("SUCCESS: All key financial metrics, footnotes, numbers, names, and regulatory terms verified in document!")
